"""Helpers for building small .docx fixtures programmatically with
python-docx, for the extractor/pipeline test suite.

Kept out of conftest.py so it can be imported directly by tests that
want to build custom variants (e.g. a docx with an explicit language)
without going through a fixture.
"""

from __future__ import annotations

import io
import zipfile
from pathlib import Path

import docx
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from PIL import Image as PILImage

_FRENCH_PARAGRAPHS = [
    "Ceci est un document de démonstration entièrement rédigé en français.",
    "Il sert à vérifier que la détection automatique de la langue du "
    "contenu fonctionne correctement lorsque aucune métadonnée de langue "
    "n'est présente dans le fichier source.",
    "Le pipeline doit détecter le français avec une confiance suffisante "
    "et ne pas se rabattre silencieusement sur l'anglais par défaut.",
]

# A tiny in-memory 4x4 PNG, used as the fixture's "image without alt text".
def _tiny_png_bytes() -> bytes:
    img = PILImage.new("RGB", (4, 4), color=(200, 30, 30))
    buf = io.BytesIO()
    img.save(buf, format="PNG")
    return buf.getvalue()


def add_hyperlink(paragraph, url: str, text: str) -> None:
    part = paragraph.part
    r_id = part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    run.append(OxmlElement("w:rPr"))
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_linked_image(doc, url: str) -> None:
    """Add a paragraph containing an image wrapped in a hyperlink with no
    separate link text -- the common "linked logo/banner" pattern that
    mammoth emits as `<p><a href="..."><img/></a></p>`."""
    p = doc.add_paragraph()
    run = p.add_run()
    buf = io.BytesIO(_tiny_png_bytes())
    run.add_picture(buf, width=Inches(0.5))

    part = p.part
    r_id = part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    r_el = run._r
    r_el.getparent().remove(r_el)
    hyperlink.append(r_el)
    p._p.append(hyperlink)


def set_image_alt(run, alt_text: str) -> None:
    """Set the alt text (docPr descr/name) on an inline picture run."""
    inline = run._r.find(qn("w:drawing")).find(qn("wp:inline"))
    doc_pr = inline.find(qn("wp:docPr"))
    doc_pr.set("descr", alt_text)
    doc_pr.set("name", alt_text)


def add_merged_cell_table(doc) -> None:
    """A 3x3 table with a merged header cell (colspan) and a merged data
    cell spanning two rows (rowspan)."""
    table = doc.add_table(rows=3, cols=3)
    mark_header_row(table)
    table.cell(0, 0).text = "Merged Header"
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(0, 2).text = "Count"
    table.rows[1].cells[0].text = "North"
    table.rows[1].cells[2].text = "12"
    table.rows[2].cells[0].text = "South"
    table.rows[2].cells[2].text = "9"
    table.cell(1, 1).merge(table.cell(2, 1))  # rowspan


def mark_header_row(table) -> None:
    """Set Word's 'repeat as header row' property on a table's first row."""
    tr = table.rows[0]._tr
    trPr = OxmlElement("w:trPr")
    trPr.append(OxmlElement("w:tblHeader"))
    tr.insert(0, trPr)


def add_nested_bullet_list(doc, top_text: str, nested_text: str) -> None:
    """Add a two-level nested bullet list using an explicit numId/ilvl,
    backed by a real multi-level numbering definition injected into the
    package after save() (python-docx's default template's built-in
    "List Bullet"/"List Bullet 2" styles use *separate* single-level
    numId's, not a shared multi-level list, so they don't nest under
    mammoth -- this fixture builds the real thing instead)."""
    num_id = 900

    def add_item(text: str, ilvl: int):
        p = doc.add_paragraph(text, style="List Bullet")
        pPr = p._p.get_or_add_pPr()
        numPr = OxmlElement("w:numPr")
        ilvl_el = OxmlElement("w:ilvl")
        ilvl_el.set(qn("w:val"), str(ilvl))
        numId_el = OxmlElement("w:numId")
        numId_el.set(qn("w:val"), str(num_id))
        numPr.append(ilvl_el)
        numPr.append(numId_el)
        pPr.append(numPr)

    add_item(top_text, 0)
    add_item(nested_text, 1)


_NUMBERING_XML_TEMPLATE = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:numbering xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
<w:abstractNum w:abstractNumId="900">
<w:multiLevelType w:val="hybridMultilevel"/>
<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="bullet"/><w:lvlText w:val="&#8226;"/>
<w:lvlJc w:val="left"/><w:pPr><w:ind w:left="720" w:hanging="360"/></w:pPr></w:lvl>
<w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="bullet"/><w:lvlText w:val="o"/>
<w:lvlJc w:val="left"/><w:pPr><w:ind w:left="1440" w:hanging="360"/></w:pPr></w:lvl>
</w:abstractNum>
<w:num w:numId="900"><w:abstractNumId w:val="900"/></w:num>
</w:numbering>"""


def _inject_numbering_part(docx_bytes: bytes) -> bytes:
    """Post-process a saved .docx's zip to add a real multi-level
    numbering.xml part (see add_nested_bullet_list docstring)."""
    zin = zipfile.ZipFile(io.BytesIO(docx_bytes), "r")
    names = zin.namelist()
    has_numbering = "word/numbering.xml" in names

    outbuf = io.BytesIO()
    zout = zipfile.ZipFile(outbuf, "w", zipfile.ZIP_DEFLATED)
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename == "word/numbering.xml":
            data = _NUMBERING_XML_TEMPLATE.encode("utf-8")
        elif item.filename == "[Content_Types].xml" and not has_numbering:
            data = data.replace(
                b"</Types>",
                b'<Override PartName="/word/numbering.xml" '
                b'ContentType="application/vnd.openxmlformats-officedocument'
                b'.wordprocessingml.numbering+xml"/></Types>',
            )
        elif item.filename == "word/_rels/document.xml.rels" and not has_numbering:
            data = data.replace(
                b"</Relationships>",
                b'<Relationship Id="rIdFixtureNumbering900" '
                b'Type="http://schemas.openxmlformats.org/officeDocument/2006'
                b'/relationships/numbering" Target="numbering.xml"/></Relationships>',
            )
        zout.writestr(item, data)
    if not has_numbering:
        zout.writestr("word/numbering.xml", _NUMBERING_XML_TEMPLATE)
    zout.close()
    return outbuf.getvalue()


def build_fixture_docx(path: str | Path, set_language: str | None = None) -> Path:
    """Build a .docx covering every extraction case the test suite needs:
    a heading hierarchy with a skipped level (H1 -> H3), a nested bullet
    list, a table with an explicit header row, an image with no alt
    text, and a "click here" generic link. Language is left unset unless
    `set_language` is given.
    """
    path = Path(path)
    doc = docx.Document()

    if set_language:
        doc.core_properties.language = set_language

    doc.add_heading("Fixture Document", level=1)
    doc.add_heading("Deeply Nested Section", level=3)  # heading-skip case

    p = doc.add_paragraph("This is intro text with a ")
    run = p.add_run("bold word")
    run.bold = True
    p.add_run(".")

    link_para = doc.add_paragraph("For details, ")
    add_hyperlink(link_para, "https://example.org/details", "click here")

    add_nested_bullet_list(doc, "Top-level item", "Nested item")

    table = doc.add_table(rows=2, cols=2)
    mark_header_row(table)
    table.rows[0].cells[0].text = "Region"
    table.rows[0].cells[1].text = "Count"
    table.rows[1].cells[0].text = "North"
    table.rows[1].cells[1].text = "12"

    img_buf = io.BytesIO(_tiny_png_bytes())
    doc.add_picture(img_buf, width=Inches(0.5))  # no alt text set -> MISSING_ALT

    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = _inject_numbering_part(buf.getvalue())

    path.write_bytes(docx_bytes)
    return path


def build_minimal_docx(path: str | Path) -> Path:
    """A docx with no headings at all -- for the title-promotion case."""
    path = Path(path)
    doc = docx.Document()
    doc.add_paragraph("Just a paragraph, no headings.")
    buf = io.BytesIO()
    doc.save(buf)
    path.write_bytes(buf.getvalue())
    return path


def build_docx_with_linked_image(path: str | Path) -> Path:
    """A docx whose only body content is an image wrapped in a
    hyperlink, with no separate link text (regression fixture for the
    "image inside a hyperlink silently dropped" bug)."""
    path = Path(path)
    doc = docx.Document()
    doc.add_heading("Linked Image Fixture", level=1)
    add_linked_image(doc, "https://example.org/home")
    buf = io.BytesIO()
    doc.save(buf)
    path.write_bytes(buf.getvalue())
    return path


def build_docx_with_filename_alt(path: str | Path) -> Path:
    """A docx with an image whose alt text is just its own filename
    ("Picture1.jpg") -- regression fixture for filename-like alt
    shipping verbatim instead of being treated as missing."""
    path = Path(path)
    doc = docx.Document()
    doc.add_heading("Filename Alt Fixture", level=1)
    p = doc.add_paragraph()
    run = p.add_run()
    buf_img = io.BytesIO(_tiny_png_bytes())
    run.add_picture(buf_img, width=Inches(0.5))
    set_image_alt(run, "Picture1.jpg")
    buf = io.BytesIO()
    doc.save(buf)
    path.write_bytes(buf.getvalue())
    return path


def build_docx_with_merged_table(path: str | Path) -> Path:
    """A docx with a table containing merged cells (colspan + rowspan) --
    regression fixture for the "merged cells silently corrupt the grid"
    bug."""
    path = Path(path)
    doc = docx.Document()
    doc.add_heading("Merged Table Fixture", level=1)
    add_merged_cell_table(doc)
    buf = io.BytesIO()
    doc.save(buf)
    path.write_bytes(buf.getvalue())
    return path


def build_french_text_docx(path: str | Path) -> Path:
    """A docx with real French prose and no dc:language metadata --
    regression fixture for content-based language detection."""
    path = Path(path)
    doc = docx.Document()
    doc.add_heading("Document en francais", level=1)
    for paragraph in _FRENCH_PARAGRAPHS:
        doc.add_paragraph(paragraph)
    buf = io.BytesIO()
    doc.save(buf)
    path.write_bytes(buf.getvalue())
    return path


def build_near_empty_docx(path: str | Path) -> Path:
    """A docx with almost no text -- too little for language detection
    to work confidently, so the pipeline must fall back to 'en' with a
    LANG_ASSUMED flag rather than guess."""
    path = Path(path)
    doc = docx.Document()
    doc.add_paragraph("Hi.")
    buf = io.BytesIO()
    doc.save(buf)
    path.write_bytes(buf.getvalue())
    return path
